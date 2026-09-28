#!/usr/bin/env python3
"""제품 리스트 엑셀(S~Z열)의 용지 배치값을 label_specs 반영용 SQL로 변환한다.

엑셀 열 구성(4행부터 데이터)
  C 품번 / J 제품규격(A4 등) / K 라벨수(칸) / L 표준치수 / M Spec(mm) / R 재질
  S 마진 상단 / T 마진 좌측 / U 가로 개수(열수) / V 세로 개수(행수)
  W 라벨 상하 간격 / X 라벨 좌우 간격 / Y 모서리 R / Z 라벨 바탕색 RGB

label_specs 행은 ShopProductImportService::ensureSpec() 이 만들 때와 같은 키
(width_mm, height_mm, material, labels_per_sheet)로 맞춘다. 품번 여러 개가
같은 규격을 공유하므로 규격 단위로 값을 합쳐서 내보낸다.
"""
from __future__ import annotations

import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path.home() / "Desktop" / "제품 리스트.xlsx"
OUT_SQL = ROOT.parents[0] / ".dbSchema" / "updateSQL_260923_2.sql"

COL = {
    "sku": 3, "paper_size": 10, "labels": 11, "std_size": 12, "spec_mm": 13,
    "material": 18, "product_name": 24,
    "top_margin": 19, "left_margin": 20, "columns": 21, "rows": 22,
    "v_gap": 23, "h_gap": 24, "radius": 25, "rgb": 26,
}

PAPER_MM = {
    "A3": (297.0, 420.0), "A4": (210.0, 297.0), "A5": (148.0, 210.0), "A6": (105.0, 148.0),
    "B4": (257.0, 364.0), "B5": (182.0, 257.0), "LETTER": (215.9, 279.4),
}

BLANK_TOKENS = {"", "-", "없음", "없슴", "N/A", "NA", "n/a"}


def to_float(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if text in BLANK_TOKENS:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def to_text(value) -> str:
    return "" if value is None else str(value).strip()


def parse_size(*texts) -> tuple[float | None, float | None]:
    for text in texts:
        if not text:
            continue
        m = re.search(r"([\d.]+)\s*[xX×]\s*([\d.]+)", str(text))
        if m:
            return round(float(m.group(1)), 2), round(float(m.group(2)), 2)
    return None, None


def parse_rgb(value) -> str | None:
    """'R255 G255 B255' 형태를 #RRGGBB 로. 엑셀에 R/G/B 표기 오타가 있어 숫자 순서만 본다."""
    nums = re.findall(r"\d+", to_text(value))
    if len(nums) < 3:
        return None
    r, g, b = (max(0, min(255, int(n))) for n in nums[:3])
    return f"#{r:02X}{g:02X}{b:02X}"


def sql_str(value) -> str:
    if value is None or value == "":
        return "NULL"
    escaped = str(value).replace("\\", "\\\\").replace("'", "''")
    return f"'{escaped}'"


def sql_num(value, digits: int = 3) -> str:
    if value is None:
        return "NULL"
    return f"{round(float(value), digits):.{digits}f}"


def sql_int(value) -> str:
    return "NULL" if value is None else str(int(value))


class SpecGeometry:
    """같은 규격 키를 공유하는 제품 행들을 모아 대표값 하나를 고른다."""

    def __init__(self, key: tuple) -> None:
        self.key = key
        self.rows: list[dict] = []

    def add(self, row: dict) -> None:
        self.rows.append(row)

    def fits_paper(self, row: dict) -> bool:
        width, height = self.key[0], self.key[1]
        page = PAPER_MM.get(row["paper_size"].upper())
        cols, rows_n = row["columns"], row["rows"]
        if not page or not cols or not rows_n:
            return False
        used_w = (row["left_margin"] or 0) + cols * width + (cols - 1) * (row["h_gap"] or 0)
        used_h = (row["top_margin"] or 0) + rows_n * height + (rows_n - 1) * (row["v_gap"] or 0)
        return used_w <= page[0] + 0.6 and used_h <= page[1] + 0.6

    def count_matches(self, row: dict) -> bool:
        labels = self.key[3]
        if not labels or not row["columns"] or not row["rows"]:
            return False
        return int(row["columns"]) * int(row["rows"]) == labels

    def has_grid(self, row: dict) -> bool:
        return bool(row["columns"]) and bool(row["rows"])

    def best(self) -> tuple[dict, list[str]]:
        """칸수 일치 > 용지 안에 들어맞음 > 다수결 순으로 대표 행을 고른다."""
        signature = Counter(self.signature(r) for r in self.rows)

        def score(row: dict) -> tuple:
            return (
                2 if self.count_matches(row) else 0,
                1 if self.fits_paper(row) else 0,
                signature[self.signature(row)],
                -row["excel_row"],
            )

        chosen = max(self.rows, key=score)
        warnings: list[str] = []
        if len(signature) > 1:
            others = sorted({r["sku"] for r in self.rows if self.signature(r) != self.signature(chosen)})
            warnings.append(f"값 불일치 → {chosen['sku']} 기준 채택 (다른 값: {', '.join(others)})")
        if self.key[3] and self.has_grid(chosen) and not self.count_matches(chosen):
            warnings.append(
                f"칸수({self.key[3]}) != 열수x행수({int(chosen['columns'])}x{int(chosen['rows'])})"
            )
        if self.has_grid(chosen) and not self.fits_paper(chosen):
            warnings.append("여백·간격을 더하면 용지 밖으로 넘침")
        if not self.has_grid(chosen) and self.key[3] != 1:
            warnings.append("엑셀에 행·열 값이 없어 비워 둠")
        return chosen, warnings

    @staticmethod
    def signature(row: dict) -> tuple:
        return (
            row["top_margin"], row["left_margin"], row["columns"], row["rows"],
            row["v_gap"], row["h_gap"], row["radius"], row["label_color"], row["paper_size"],
        )


def collect(xlsx: Path) -> tuple[dict, list[str], list[str]]:
    ws = openpyxl.load_workbook(xlsx, data_only=True).active
    groups: dict[tuple, SpecGeometry] = {}
    skipped: list[str] = []

    for r in range(4, ws.max_row + 1):
        sku = to_text(ws.cell(r, COL["sku"]).value)
        product_name = to_text(ws.cell(r, COL["product_name"]).value)
        if not sku and not product_name:
            continue

        width, height = parse_size(ws.cell(r, COL["spec_mm"]).value, ws.cell(r, COL["std_size"]).value)
        if not width or not height:
            skipped.append(f"{r}행 {sku or product_name}: Spec(mm)/표준치수 없음")
            continue

        labels = to_float(ws.cell(r, COL["labels"]).value)
        key = (
            width,
            height,
            to_text(ws.cell(r, COL["material"]).value) or "라벨지",
            int(labels) if labels else None,
        )

        row = {
            "excel_row": r,
            "sku": sku or product_name,
            "paper_size": to_text(ws.cell(r, COL["paper_size"]).value),
            "top_margin": to_float(ws.cell(r, COL["top_margin"]).value),
            "left_margin": to_float(ws.cell(r, COL["left_margin"]).value),
            "columns": to_float(ws.cell(r, COL["columns"]).value),
            "rows": to_float(ws.cell(r, COL["rows"]).value),
            "v_gap": to_float(ws.cell(r, COL["v_gap"]).value),
            "h_gap": to_float(ws.cell(r, COL["h_gap"]).value),
            "radius": to_float(ws.cell(r, COL["radius"]).value),
            "label_color": parse_rgb(ws.cell(r, COL["rgb"]).value),
        }
        groups.setdefault(key, SpecGeometry(key)).add(row)

    return groups, skipped, []


def build_sql(groups: dict, skipped: list[str]) -> str:
    values: list[str] = []
    notes: list[str] = []
    empty: list[str] = []

    for key in sorted(groups, key=lambda k: (k[2], k[0], k[1], k[3] or 0)):
        spec = groups[key]
        row, warnings = spec.best()
        has_layout = any(
            row[field] is not None
            for field in ("top_margin", "left_margin", "columns", "rows", "v_gap", "h_gap", "radius")
        )
        if not has_layout and not row["label_color"]:
            empty.append(f"{key[0]}x{key[1]} {key[2]} {key[3]}칸 ({row['sku']})")
            continue

        width, height, material, labels = key
        paper = row["paper_size"].upper() if row["paper_size"] else None
        # 엑셀이 0 또는 공란이면 배치를 모른다는 뜻이라 비워 둔다. 1칸짜리만 1x1로 확정할 수 있다.
        cols, rows_n = row["columns"], row["rows"]
        if not cols or not rows_n:
            cols = rows_n = 1 if labels == 1 else None
        values.append(
            "({w}, {h}, {m}, {l}, {ps}, {tm}, {lm}, {cols}, {rows}, {hg}, {vg}, {rx}, {ry}, {color}, {sku})".format(
                w=f"{width:.2f}", h=f"{height:.2f}", m=sql_str(material), l=sql_int(labels),
                ps=sql_str(paper), tm=sql_num(row["top_margin"]), lm=sql_num(row["left_margin"]),
                cols=sql_int(cols), rows=sql_int(rows_n),
                hg=sql_num(row["h_gap"]), vg=sql_num(row["v_gap"]),
                rx=sql_num(row["radius"]), ry=sql_num(row["radius"]),
                color=sql_str(row["label_color"]), sku=sql_str(row["sku"]),
            )
        )
        for warning in warnings:
            notes.append(f"{width}x{height} {material} {labels}칸 — {warning}")

    head = [
        "-- updateSQL_260923_2.sql",
        "-- 제품 리스트 엑셀(S~Z열)의 용지 배치값을 label_specs에 채운다.",
        "-- 매칭 키: width_mm, height_mm, material, labels_per_sheet",
        "--   (ShopProductImportService::ensureSpec 이 규격을 만들 때 쓰는 키와 같다)",
        "-- 엑셀 U열은 가로 개수(열수), V열은 세로 개수(행수)로 읽었다.",
        "--   헤더에는 U=행, V=열로 적혀 있지만 실제 값은 반대다. 예: A304(12칸, 63.5x72mm)는 U=3, V=4이고",
        "--   63.5x3=190.5mm(가로), 72x4=288mm(세로)라야 A4에 들어간다.",
        f"-- 반영 대상 규격: {len(values)}건",
        "",
    ]
    if notes:
        head.append("-- [검토 필요] 엑셀 값이 서로 어긋나는 규격")
        head.extend(f"--   {n}" for n in notes)
        head.append("")
    if empty:
        head.append("-- [건너뜀] 배치값이 비어 있는 규격")
        head.extend(f"--   {n}" for n in empty)
        head.append("")
    if skipped:
        head.append("-- [건너뜀] 치수를 읽지 못한 엑셀 행")
        head.extend(f"--   {n}" for n in skipped)
        head.append("")

    body = """DROP TABLE IF EXISTS tmp_label_spec_geometry;

CREATE TABLE tmp_label_spec_geometry (
    width_mm DECIMAL(8,2) NOT NULL,
    height_mm DECIMAL(8,2) NOT NULL,
    material VARCHAR(80) NOT NULL,
    labels_per_sheet INT UNSIGNED NULL,
    paper_size VARCHAR(20) NULL,
    top_margin_mm DECIMAL(7,3) NULL,
    left_margin_mm DECIMAL(7,3) NULL,
    columns_count SMALLINT UNSIGNED NULL,
    rows_count SMALLINT UNSIGNED NULL,
    h_gap_mm DECIMAL(7,3) NULL,
    v_gap_mm DECIMAL(7,3) NULL,
    corner_radius_x_mm DECIMAL(7,3) NULL,
    corner_radius_y_mm DECIMAL(7,3) NULL,
    label_color VARCHAR(7) NULL,
    sample_sku VARCHAR(80) NULL,
    KEY idx_tmp_spec_key (width_mm, height_mm, material, labels_per_sheet)
) ENGINE=InnoDB DEFAULT CHARSET=utf8 COLLATE=utf8_general_ci;

INSERT INTO tmp_label_spec_geometry
    (width_mm, height_mm, material, labels_per_sheet, paper_size,
     top_margin_mm, left_margin_mm, columns_count, rows_count,
     h_gap_mm, v_gap_mm, corner_radius_x_mm, corner_radius_y_mm, label_color, sample_sku)
VALUES
{values};

UPDATE label_specs s
JOIN tmp_label_spec_geometry t
  ON s.width_mm = t.width_mm
 AND s.height_mm = t.height_mm
 AND s.material = t.material
 AND s.labels_per_sheet <=> t.labels_per_sheet
SET s.paper_size         = t.paper_size,
    s.top_margin_mm      = t.top_margin_mm,
    s.left_margin_mm     = t.left_margin_mm,
    s.columns_count      = t.columns_count,
    s.rows_count         = t.rows_count,
    s.h_gap_mm           = t.h_gap_mm,
    s.v_gap_mm           = t.v_gap_mm,
    s.corner_radius_x_mm = t.corner_radius_x_mm,
    s.corner_radius_y_mm = t.corner_radius_y_mm,
    s.label_color        = t.label_color,
    s.updated_at         = NOW();

-- 반영되지 않은 엑셀 규격 확인 (결과가 비어 있어야 정상)
SELECT t.sample_sku, t.width_mm, t.height_mm, t.material, t.labels_per_sheet
FROM tmp_label_spec_geometry t
LEFT JOIN label_specs s
       ON s.width_mm = t.width_mm
      AND s.height_mm = t.height_mm
      AND s.material = t.material
      AND s.labels_per_sheet <=> t.labels_per_sheet
WHERE s.id IS NULL;

DROP TABLE tmp_label_spec_geometry;
""".replace("{values}", ",\n".join(values))

    return "\n".join(head) + "\n" + body


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx.exists():
        print(f"엑셀을 찾지 못했습니다: {xlsx}", file=sys.stderr)
        return 1

    groups, skipped, _ = collect(xlsx)
    sql = build_sql(groups, skipped)
    OUT_SQL.parent.mkdir(parents=True, exist_ok=True)
    OUT_SQL.write_text(sql, encoding="utf-8")
    print(f"규격 {len(groups)}건 분석 → {OUT_SQL}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
